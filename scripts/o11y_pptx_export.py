"""
Build a concise PowerPoint executive deck from consolidated health-check JSON.

Optional dependency: ``python-pptx`` (see ``requirements-presentation.txt``).
When a Splunk template ``.pptx`` is provided (or chosen by the orchestrator), the deck **keeps the
template’s slide masters and layouts** (fonts, colors, branding) but **removes every existing slide**
in that file so the export contains only generated health-check slides. The runner defaults to the
**dark** bundled master unless ``--pptx-light`` / profile theme selects the light file.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from o11y_presentation_util import (
    color_rank,
    extract_executive_bullets_from_markdown,
    strip_inline_markdown,
)

logger = logging.getLogger(__name__)


def _deck_include(combined: dict[str, Any], key: str) -> bool:
    """Whether the executive deck should include slides for ``key`` (new runs set ``reportSections``)."""
    rs = combined.get("reportSections")
    if not isinstance(rs, dict):
        return True
    return bool(rs.get(key, True))


def _severity_sort_key(row: dict[str, Any]) -> tuple[int, str]:
    sev = row.get("color") or row.get("severity")
    return (color_rank(sev), str(row).lower())


def _top_rows(rows: list[dict[str, Any]], *, max_n: int) -> list[dict[str, Any]]:
    if not rows or max_n <= 0:
        return []
    return sorted(rows, key=_severity_sort_key)[:max_n]


def _truncate(s: Any, max_len: int = 54) -> str:
    t = str(s if s is not None else "").replace("\n", " ").strip()
    if len(t) <= max_len:
        return t
    return t[: max_len - 1] + "…"


def _license_table_rows(combined: dict[str, Any], *, max_n: int) -> tuple[list[str], list[list[str]]]:
    lic = combined.get("license") or {}
    rows = lic.get("rows") or []
    hdr = ["Product / family", "Severity", "Util %"]
    data: list[list[str]] = []
    for r in _top_rows(rows, max_n=max_n):
        prod = _truncate(r.get("productFamily") or r.get("product") or r.get("name") or "—")
        sev = _truncate(r.get("severity") or "—")
        pct = r.get("utilization_pct")
        pct_s = "—" if pct is None else f"{float(pct):.1f}%"
        data.append([prod, sev, pct_s])
    return hdr, data


def _apm_summary_bullets(combined: dict[str, Any], *, max_checks: int = 8) -> list[str]:
    apm = combined.get("apm") or {}
    if apm.get("error"):
        return ["APM: automation did not complete — see markdown / PDF."]
    if not apm:
        return ["APM: not executed in this run."]
    wh = apm.get("window_hours", "—")
    lines = [f"Snapshot window: {wh} hour(s)."]
    checks = apm.get("checks") or {}
    if not isinstance(checks, dict):
        return lines
    n = 0
    for key, blk in checks.items():
        if n >= max_checks:
            lines.append("… additional APM checks — see detailed PDF / markdown report.")
            break
        if not isinstance(blk, dict):
            continue
        rows = blk.get("rows") or []
        red = sum(1 for x in rows if str(x.get("color", "")).lower() == "red")
        yel = sum(1 for x in rows if str(x.get("color", "")).lower() == "yellow")
        title = str(key).replace("_", " ").strip().title()
        lines.append(f"{title}: {len(rows)} row(s); Red {red}, Yellow {yel}.")
        n += 1
    return lines


def _detector_sample_table(combined: dict[str, Any], *, max_n: int) -> tuple[list[str], list[list[str]]]:
    det = combined.get("detectors") or {}
    hdr = ["Check", "Detector", "Color"]
    data: list[list[str]] = []
    if not det or det.get("error"):
        return hdr, data
    checks = det.get("checks") or {}
    flat: list[tuple[str, dict[str, Any]]] = []
    for cname, blk in checks.items():
        if not isinstance(blk, dict):
            continue
        for r in blk.get("rows") or []:
            if isinstance(r, dict):
                flat.append((str(cname), r))
    flat.sort(key=lambda t: _severity_sort_key(t[1]))
    for cname, r in flat[:max_n]:
        dname = _truncate(r.get("detectorName") or r.get("name") or "—", 40)
        col = _truncate(r.get("color") or "—", 8)
        data.append([_truncate(cname, 28), dname, col])
    return hdr, data


def _dashboard_sample_table(combined: dict[str, Any], *, max_n: int) -> tuple[list[str], list[list[str]]]:
    dash = combined.get("dashboards") or {}
    hdr = ["Check", "Dashboard / detail", "Color"]
    data: list[list[str]] = []
    if not dash or dash.get("error"):
        return hdr, data
    checks = dash.get("checks") or {}
    flat: list[tuple[str, dict[str, Any]]] = []
    for cname, blk in checks.items():
        if not isinstance(blk, dict):
            continue
        for r in blk.get("rows") or []:
            if isinstance(r, dict):
                flat.append((str(cname), r))
    flat.sort(key=lambda t: _severity_sort_key(t[1]))
    for cname, r in flat[:max_n]:
        name = _truncate(
            r.get("dashboardName") or r.get("chartTitle") or r.get("name") or "—",
            42,
        )
        col = _truncate(r.get("color") or "—", 8)
        data.append([_truncate(cname, 28), name, col])
    return hdr, data


def _token_sample_table(combined: dict[str, Any], *, max_n: int) -> tuple[list[str], list[list[str]]]:
    tok = combined.get("tokens") or {}
    hdr = ["Token name", "Type", "Expiry / note"]
    data: list[list[str]] = []
    if not tok or tok.get("error"):
        return hdr, data
    rows: list[dict[str, Any]] = []
    for key in ("expiredTokens", "nearExpiration"):
        blk = (tok.get("checks") or {}).get(key) or {}
        for r in blk.get("rows") or []:
            if isinstance(r, dict):
                rows.append(r)
    for r in rows[:max_n]:
        tt = r.get("tokenTypes") or r.get("tokenType")
        if isinstance(tt, list):
            tt_s = ", ".join(str(x) for x in tt if x)
        else:
            tt_s = tt
        data.append(
            [
                _truncate(r.get("name") or r.get("tokenName") or "—", 36),
                _truncate(tt_s or "—", 14),
                _truncate(r.get("expiredDate") or r.get("expirationDate") or "—", 16),
            ]
        )
    return hdr, data


def _otel_sample_table(combined: dict[str, Any], *, max_n: int) -> tuple[list[str], list[list[str]]]:
    otel = combined.get("otelCollectors") or {}
    hdr = ["Host", "Color", "Version"]
    data: list[list[str]] = []
    if not otel or otel.get("error"):
        return hdr, data
    rows = ((otel.get("checks") or {}).get("collectorsByVersion") or {}).get("rows") or []
    for r in _top_rows(rows, max_n=max_n):
        data.append(
            [
                _truncate(r.get("hostName") or "—", 28),
                _truncate(r.get("color") or "—", 8),
                _truncate(r.get("version") or "—", 12),
            ]
        )
    return hdr, data


def _im_metrics_top(combined: dict[str, Any], *, max_n: int) -> tuple[list[str], list[list[str]]]:
    im = combined.get("im") or {}
    hdr = ["Metric", "Avg MTS / hour"]
    data: list[list[str]] = []
    metrics = im.get("metrics") or []
    if not metrics:
        return hdr, data
    sorted_m = sorted(
        metrics,
        key=lambda m: float(m.get("averageHourlyMts") or m.get("avgMtsPerHour") or m.get("mts") or 0),
        reverse=True,
    )
    for m in sorted_m[:max_n]:
        name = _truncate(m.get("metricName") or m.get("metric") or m.get("name") or "—", 44)
        v = m.get("averageHourlyMts") or m.get("avgMtsPerHour") or m.get("mts")
        vs = "—" if v is None else _truncate(str(v), 12)
        data.append([name, vs])
    return hdr, data


def _strip_template_slides(prs: Any) -> None:
    """Remove all slides from a loaded presentation; keep masters/layouts for theming.

    Bundled Splunk templates include many reference/branding slides that must not ship in the
    customer-facing executive deck.
    """
    sld_id_lst = prs.slides._sldIdLst  # CT_SlideIdList — same object python-pptx uses internally
    for sld_id in list(sld_id_lst):
        r_id = sld_id.rId
        prs.part.drop_rel(r_id)
        parent = sld_id.getparent()
        if parent is not None:
            parent.remove(sld_id)
    if len(prs.slides):
        logger.warning(
            "PPTX: %s template slide id(s) still present after strip — deck may include extra slides",
            len(prs.slides),
        )


def _blank_or_title_content_layout(prs: Any):
    """Prefer TITLE_AND_CONTENT; fall back to layout index 1 or 0."""
    for layout in prs.slide_layouts:
        try:
            name = (layout.name or "").upper()
        except Exception:
            name = ""
        if "TITLE_AND_CONTENT" in name.replace(" ", "_") or "TITLE AND CONTENT" in name:
            return layout
    if len(prs.slide_layouts) > 1:
        return prs.slide_layouts[1]
    return prs.slide_layouts[0]


def _set_slide_title(slide: Any, title: str) -> None:
    """Set title placeholder if present; otherwise add a textbox heading."""
    from pptx.util import Inches, Pt  # type: ignore[import-untyped]

    t = strip_inline_markdown(title)
    if getattr(slide.shapes, "title", None) is not None:
        slide.shapes.title.text = t
        return
    box = slide.shapes.add_textbox(Inches(0.42), Inches(0.22), Inches(9.15), Inches(0.85))
    tf = box.text_frame
    tf.text = t
    for p in tf.paragraphs:
        p.font.size = Pt(22)
        p.font.bold = True


def _place_subtitle_on_slide(slide: Any, sub_text: str) -> None:
    """Splunk/custom masters often omit standard placeholder idx 1; fall back to a textbox."""
    from pptx.util import Inches, Pt  # type: ignore[import-untyped]

    if not sub_text:
        return
    try:
        slide.placeholders[1].text = sub_text
        return
    except KeyError:
        pass
    try:
        from pptx.enum.shapes import PP_PLACEHOLDER  # type: ignore[import-untyped]

        for shape in slide.placeholders:
            try:
                if shape.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                    shape.text = sub_text
                    return
            except Exception:
                continue
    except Exception:
        pass
    box = slide.shapes.add_textbox(Inches(0.42), Inches(1.05), Inches(9.15), Inches(1.45))
    tf = box.text_frame
    tf.text = sub_text
    for p in tf.paragraphs:
        p.font.size = Pt(13)


def _body_text_frame(slide: Any) -> Any:
    """Title+content layouts usually use placeholder 1 for body; templates vary."""
    from pptx.util import Inches  # type: ignore[import-untyped]

    try:
        return slide.placeholders[1].text_frame
    except KeyError:
        pass
    try:
        from pptx.enum.shapes import PP_PLACEHOLDER  # type: ignore[import-untyped]

        for shape in slide.placeholders:
            try:
                if shape.placeholder_format.type in (
                    PP_PLACEHOLDER.BODY,
                    PP_PLACEHOLDER.OBJECT,
                ):
                    return shape.text_frame
            except Exception:
                continue
    except Exception:
        pass
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.idx > 0:
                return shape.text_frame
        except Exception:
            continue
    box = slide.shapes.add_textbox(Inches(0.55), Inches(1.4), Inches(9.0), Inches(4.8))
    return box.text_frame


def _add_bullet_slide(prs: Any, title: str, bullets: list[str], *, font_pt: int = 14) -> None:
    from pptx.util import Pt  # type: ignore[import-untyped]

    layout = _blank_or_title_content_layout(prs)
    slide = prs.slides.add_slide(layout)
    _set_slide_title(slide, title)
    tf = _body_text_frame(slide)
    tf.clear()
    for i, b in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = strip_inline_markdown(b)
        p.level = 0
        p.font.size = Pt(font_pt)


def _add_title_slide(prs: Any, title: str, subtitle_lines: list[str]) -> None:
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    _set_slide_title(slide, title)
    sub_text = "\n".join(strip_inline_markdown(x) for x in subtitle_lines if x)
    _place_subtitle_on_slide(slide, sub_text)


def _add_table_slide(
    prs: Any,
    title: str,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    from pptx.util import Inches  # type: ignore[import-untyped]

    if not rows:
        _add_bullet_slide(prs, title, ["No prioritized rows — see PDF for full tables."])
        return
    try:
        blank = prs.slide_layouts[6]
    except IndexError:
        blank = _blank_or_title_content_layout(prs)
    slide = prs.slides.add_slide(blank)
    _set_slide_title(slide, title)
    r_count, c_count = len(rows) + 1, len(headers)
    left, top, width, height = Inches(0.4), Inches(1.35), Inches(9.1), min(Inches(0.42 * (r_count + 1)), Inches(4.9))
    table = slide.shapes.add_table(r_count, c_count, left, top, width, height).table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = strip_inline_markdown(h)
        for para in cell.text_frame.paragraphs:
            para.font.bold = True
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            table.cell(i + 1, j).text = strip_inline_markdown(str(val))


def export_pptx(
    *,
    combined: dict[str, Any],
    consolidated_markdown: str,
    out_path: Path,
    template_path: Path | None = None,
    max_table_rows: int = 7,
    max_exec_bullets: int = 14,
) -> str | None:
    """
    Write a ``.pptx`` file. Returns an error message string on failure, or ``None`` on success.
    """
    try:
        from pptx import Presentation  # type: ignore[import-untyped]
    except ImportError:
        return "python-pptx is not installed. Install optional deps: pip install -r requirements-presentation.txt"

    out_path = out_path.expanduser().resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        if template_path and template_path.is_file():
            prs = Presentation(str(template_path))
            n_tpl_slides = len(prs.slides)
            logger.info("PPTX: loaded template %s (%s slide(s))", template_path, n_tpl_slides)
            _strip_template_slides(prs)
            if n_tpl_slides:
                logger.info(
                    "PPTX: removed %s template/branding slide(s); deck will contain generated content only",
                    n_tpl_slides,
                )
        else:
            prs = Presentation()
            if template_path:
                logger.warning("PPTX template not found (%s); using blank presentation", template_path)
    except Exception as e:
        return f"Could not open PowerPoint template/presentation: {e}"

    title = str(combined.get("title") or "Splunk Observability Health Check")
    customer = str(combined.get("customer") or "—")
    realm = str(combined.get("realm") or "—")
    ts = str(combined.get("generated_utc") or "")
    subtitle = [customer, f"Realm: {realm}", f"Generated UTC: {ts}" if ts else ""]

    _add_title_slide(prs, title, subtitle)

    exec_bullets = extract_executive_bullets_from_markdown(consolidated_markdown)
    if not exec_bullets:
        exec_bullets = ["No executive summary section found — see markdown report."]
    else:
        note = (
            "Full tables and findings appear in the detailed PDF/markdown artifact "
            "(this deck shows highlights only)."
        )
        exec_bullets = exec_bullets[:max_exec_bullets] + ([] if note in exec_bullets else [note])
    _add_bullet_slide(prs, "Executive summary", exec_bullets)

    if _deck_include(combined, "license"):
        hdr, lic_rows = _license_table_rows(combined, max_n=max_table_rows)
        lic = combined.get("license") or {}
        lic_note = []
        lr = lic.get("rows") or []
        if lr:
            lic_note.append(f"License entitlements (table): showing top {max_table_rows} by severity emphasis.")
        else:
            lic_note.append("License: no entitlement rows.")
        _add_bullet_slide(prs, "License utilization (overview)", lic_note)
        if lic_rows:
            _add_table_slide(prs, "License — prioritized rows", hdr, lic_rows)

    if _deck_include(combined, "platformEngagement"):
        pe = combined.get("platformEngagement") or {}
        if pe.get("error"):
            _add_bullet_slide(prs, "Platform engagement", ["Automation incomplete — see report."])
        else:
            nk = len(pe.get("kpis") or [])
            finds = pe.get("findings") or []
            pb = [f"KPI series: {nk}"]
            for f in finds[:5]:
                pb.append(strip_inline_markdown(str(f))[:420])
            if len(finds) > 5:
                pb.append("… see PDF for remaining engagement findings.")
            _add_bullet_slide(prs, "Platform engagement", pb)

    if _deck_include(combined, "im"):
        im = combined.get("im") or {}
        ii = combined.get("im_integrations") or {}
        imb = []
        if im.get("error"):
            imb.append("IM metrics usage: did not complete.")
        else:
            imb.append(f"IM metrics rows: {len(im.get('metrics') or [])}.")
        if ii.get("error"):
            imb.append("IM integrations: did not complete.")
        else:
            imb.append(f"Integrations listed: {len(ii.get('integrations') or [])}.")
        _add_bullet_slide(prs, "Infrastructure monitoring", imb)
        mh, md = _im_metrics_top(combined, max_n=max_table_rows)
        if md:
            _add_table_slide(prs, "IM — highest volume metrics (sample)", mh, md)

    rs_det = _deck_include(combined, "detectors")
    rs_dash = _deck_include(combined, "dashboards")
    if rs_det or rs_dash:
        det = combined.get("detectors") or {}
        dab = combined.get("dashboards") or {}
        db_ul: list[str] = []
        if rs_det:
            if det.get("error"):
                db_ul.append("Detectors: automation failed.")
            else:
                db_ul.append(
                    f"Detectors analyzed: {det.get('detectorsAnalyzed')} "
                    f"of {det.get('detectorListTotal')} listed."
                )
        if rs_dash:
            if dab.get("error"):
                db_ul.append("Dashboards: automation failed.")
            else:
                db_ul.append(
                    f"Dashboards analyzed: {dab.get('dashboardsAnalyzed')} "
                    f"of {dab.get('dashboardListTotal')} listed."
                )
        if rs_det and rs_dash:
            db_ul.append("Sample Red/Yellow rows follow (full inventory in PDF).")
            cov_title = "Detectors & dashboards (coverage)"
        elif rs_det:
            db_ul.append("Sample detector rows follow (full inventory in PDF).")
            cov_title = "Detectors (coverage)"
        else:
            db_ul.append("Sample dashboard rows follow (full inventory in PDF).")
            cov_title = "Dashboards (coverage)"
        _add_bullet_slide(prs, cov_title, db_ul)

        if rs_det:
            dh, dd = _detector_sample_table(combined, max_n=max_table_rows)
            if dd:
                _add_table_slide(prs, "Detectors — sample high-priority rows", dh, dd)
        if rs_dash:
            bh, bd = _dashboard_sample_table(combined, max_n=max_table_rows)
            if bd:
                _add_table_slide(prs, "Dashboards — sample high-priority rows", bh, bd)

    if _deck_include(combined, "apm"):
        _add_bullet_slide(prs, "APM", _apm_summary_bullets(combined))

    rum = combined.get("rum") or {}
    syn = combined.get("synthetics") or {}
    tok = combined.get("tokens") or {}
    otel = combined.get("otelCollectors") or {}
    misc = []
    misc_labels: list[str] = []
    if _deck_include(combined, "rum"):
        misc_labels.append("RUM")
        if rum.get("error"):
            misc.append("RUM: automation failed.")
        else:
            misc.append(f"RUM: lookback {rum.get('lookbackHours', '—')}h; review volume/MMS in PDF.")
    if _deck_include(combined, "synthetics"):
        misc_labels.append("Synthetics")
        if syn.get("error"):
            misc.append("Synthetics: automation failed.")
        else:
            misc.append(f"Synthetics tests listed: {syn.get('testsAnalyzed', '—')}.")
    if _deck_include(combined, "tokens"):
        misc_labels.append("Tokens")
        if tok.get("error"):
            misc.append("Tokens: automation failed.")
        else:
            et = len(((tok.get("checks") or {}).get("expiredTokens") or {}).get("rows") or [])
            ne = len(((tok.get("checks") or {}).get("nearExpiration") or {}).get("rows") or [])
            misc.append(f"Tokens: expired {et}; near-expiration {ne} (detail in PDF).")
    if _deck_include(combined, "otelCollectors"):
        misc_labels.append("OTel collectors")
        if otel.get("error"):
            misc.append("OpenTelemetry collectors: automation failed.")
        else:
            n_ot = len(((otel.get("checks") or {}).get("collectorsByVersion") or {}).get("rows") or [])
            misc.append(f"OpenTelemetry collectors: {n_ot} instance row(s).")
    if misc:
        misc.append("")
        misc.append("Next steps: align on critical Red/Yellow items; use PDF appendix for exhaustive tables.")
        slide_title = " · ".join(misc_labels) if misc_labels else "Product modules"
        _add_bullet_slide(prs, slide_title, misc)

    if _deck_include(combined, "tokens"):
        th, td = _token_sample_table(combined, max_n=max_table_rows)
        if td:
            _add_table_slide(prs, "Tokens — expired / near-expiration (sample)", th, td)
    if _deck_include(combined, "otelCollectors"):
        oh, od = _otel_sample_table(combined, max_n=max_table_rows)
        if od:
            _add_table_slide(prs, "OpenTelemetry collectors (sample)", oh, od)

    try:
        prs.save(str(out_path))
    except Exception as e:
        return f"Could not save PowerPoint file: {e}"
    logger.info("Wrote PowerPoint deck: %s (%s slides)", out_path, len(prs.slides))
    return None
